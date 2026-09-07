"""Generates the synthetic resume fixtures under tests/fixtures/.

ALL fixtures produced by this script are synthetic test data: fake names,
fake companies, fake universities, fake contact details. None of it
describes a real person.

Not part of the project's runtime or test dependencies -- it needs two
packages that the shipped app never imports:

    pip install reportlab pillow

Run from the repository root:

    python3 scripts/generate_test_fixtures.py

Fixtures are checked into git as static binary files so the test suite
itself does not depend on reportlab/pillow being installed; this script
exists purely so the fixtures are reproducible/regeneratable.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import docx
from PIL import Image
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

FIXTURES_DIR = Path(__file__).resolve().parent.parent / "tests" / "fixtures"


def _pdf_canvas(path: Path) -> canvas.Canvas:
    return canvas.Canvas(str(path), pagesize=letter)


# --- 1. normal_resume.pdf --------------------------------------------------
def make_normal_resume_pdf() -> None:
    path = FIXTURES_DIR / "normal_resume.pdf"
    c = _pdf_canvas(path)
    y = 750

    def heading(text: str) -> None:
        nonlocal y
        y -= 28
        c.setFont("Helvetica-Bold", 14)
        c.drawString(72, y, text)

    def body(text: str) -> None:
        nonlocal y
        y -= 18
        c.setFont("Helvetica", 10)
        c.drawString(72, y, text)

    c.setFont("Helvetica-Bold", 18)
    c.drawString(72, y, "Jordan Testcase")
    y -= 18
    c.setFont("Helvetica", 10)
    c.drawString(72, y, "jordan.testcase@example.com  |  (555) 123-4567")

    heading("Experience")
    body("Senior Software Engineer, Fixture Industries Inc., 2021 - Present")
    body("Built internal tooling for synthetic test data generation.")
    body("Software Engineer, Example Corp, 2018 - 2021")
    body("Worked on backend services.")

    heading("Education")
    body("Bachelor of Science in Computer Science, Sample State University, 2018")

    heading("Skills")
    body("Python, FastAPI, SQL, Docker, Git")

    c.save()


# --- 2. normal_resume.docx --------------------------------------------------
def make_normal_resume_docx() -> None:
    path = FIXTURES_DIR / "normal_resume.docx"
    d = docx.Document()
    d.add_paragraph("Jordan Testcase", style="Title")
    d.add_paragraph("jordan.testcase@example.com | (555) 123-4567")

    d.add_heading("Experience", level=1)
    d.add_paragraph("Senior Software Engineer, Fixture Industries Inc., 2021 - Present")
    d.add_paragraph("Built internal tooling for synthetic test data generation.")
    d.add_paragraph("Software Engineer, Example Corp, 2018 - 2021")

    d.add_heading("Education", level=1)
    d.add_paragraph("Bachelor of Science in Computer Science, Sample State University, 2018")

    d.add_heading("Skills", level=1)
    d.add_paragraph("Python, FastAPI, SQL, Docker, Git")

    d.save(str(path))


# --- 3. table_resume.docx ---------------------------------------------------
def make_table_resume_docx() -> None:
    path = FIXTURES_DIR / "table_resume.docx"
    d = docx.Document()
    d.add_paragraph("Taylor Fixture", style="Title")
    d.add_paragraph("taylor.fixture@example.com")

    d.add_heading("Skills", level=1)
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skill"
    table.rows[0].cells[1].text = "Level"
    for skill, level in [("Python", "Expert"), ("SQL", "Advanced"), ("Docker", "Intermediate")]:
        row = table.add_row()
        row.cells[0].text = skill
        row.cells[1].text = level

    d.add_heading("Education", level=1)
    d.add_paragraph("Master of Science in Data Science, Sample State University, 2020")

    d.save(str(path))


# --- 3b. table_resume.pdf ---------------------------------------------------
def make_table_resume_pdf() -> None:
    path = FIXTURES_DIR / "table_resume.pdf"
    doc = SimpleDocTemplate(str(path))
    data = [
        ["Skill", "Level"],
        ["Python", "Expert"],
        ["SQL", "Advanced"],
        ["Docker", "Intermediate"],
    ]
    table = Table(data)
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    doc.build([table])


# --- 4. two_column_resume.pdf -----------------------------------------------
def make_two_column_resume_pdf() -> None:
    path = FIXTURES_DIR / "two_column_resume.pdf"
    c = _pdf_canvas(path)

    c.setFont("Helvetica-Bold", 18)
    c.drawString(72, 760, "Casey Sidebar")
    c.setFont("Helvetica", 10)
    c.drawString(72, 742, "casey.sidebar@example.com")

    left_x = 72
    right_x = 330
    y = 700

    c.setFont("Helvetica-Bold", 13)
    c.drawString(left_x, y, "Experience")
    c.setFont("Helvetica-Bold", 13)
    c.drawString(right_x, y, "Skills")

    c.setFont("Helvetica", 10)
    left_lines = [
        "Software Engineer, Example Corp, 2019 - Present",
        "Built distributed systems.",
        "Junior Engineer, Fixture Industries Inc., 2017 - 2019",
    ]
    right_lines = [
        "Python",
        "SQL",
        "Docker",
        "Kubernetes",
    ]
    for i, line in enumerate(left_lines):
        c.drawString(left_x, y - 20 - (i * 16), line)
    for i, line in enumerate(right_lines):
        c.drawString(right_x, y - 20 - (i * 16), line)

    c.save()


# --- 5. unusual_order_resume.pdf --------------------------------------------
def make_unusual_order_resume_pdf() -> None:
    path = FIXTURES_DIR / "unusual_order_resume.pdf"
    c = _pdf_canvas(path)
    y = 750

    def heading(text: str) -> None:
        nonlocal y
        y -= 28
        c.setFont("Helvetica-Bold", 14)
        c.drawString(72, y, text)

    def body(text: str) -> None:
        nonlocal y
        y -= 18
        c.setFont("Helvetica", 10)
        c.drawString(72, y, text)

    c.setFont("Helvetica-Bold", 18)
    c.drawString(72, y, "Morgan Reorder")
    y -= 18
    c.setFont("Helvetica", 10)
    c.drawString(72, y, "morgan.reorder@example.com")

    heading("Core Competencies")
    body("Python, Go, SQL, Terraform")

    heading("Employment History")
    body("Platform Engineer, Fixture Industries Inc., 2020 - Present")

    heading("Academic History")
    body("Master of Science in Software Engineering, Sample State University, 2016")

    c.save()


# --- 5b. headingless_resume.pdf ----------------------------------------------
def make_headingless_resume_pdf() -> None:
    """No section headings at all -- relies entirely on segmentation's
    content-based fallbacks (date range -> Experience, degree keyword ->
    Education, comma-list -> Skills) and contact-by-content detection."""
    path = FIXTURES_DIR / "headingless_resume.pdf"
    c = _pdf_canvas(path)
    y = 750

    def body(text: str) -> None:
        nonlocal y
        y -= 20
        c.setFont("Helvetica", 10)
        c.drawString(72, y, text)

    c.setFont("Helvetica", 10)
    c.drawString(72, y, "Riley Headingless")
    y -= 18
    c.drawString(72, y, "riley.headingless@example.com")

    body("Backend Engineer, Example Corp, 2020 - Present")
    body("Bachelor of Science in Computer Science, Sample State University, 2016")
    body("Python, SQL, Docker, Kubernetes, Terraform")

    c.save()


# --- 5c. varied_formats_resume.pdf -------------------------------------------
def make_varied_formats_resume_pdf() -> None:
    """For validation/ dataset coverage the other fixtures don't exercise:
    contact links, "Title at Company" phrasing, month-precision date
    ranges (both "Mon YYYY" and "MM/YYYY" forms), an abbreviated degree
    ("M.S."), and skill aliases/taxonomy gaps (see validation/README.md)."""
    path = FIXTURES_DIR / "varied_formats_resume.pdf"
    c = _pdf_canvas(path)
    y = 750

    def heading(text: str) -> None:
        nonlocal y
        y -= 28
        c.setFont("Helvetica-Bold", 14)
        c.drawString(72, y, text)

    def body(text: str) -> None:
        nonlocal y
        y -= 18
        c.setFont("Helvetica", 10)
        c.drawString(72, y, text)

    c.setFont("Helvetica-Bold", 18)
    c.drawString(72, y, "Alex Fixture")
    y -= 18
    c.setFont("Helvetica", 10)
    c.drawString(72, y, "alex.fixture@example.com  |  (555) 987-6543")
    y -= 16
    c.drawString(72, y, "linkedin.com/in/alexfixture  |  github.com/alexfixture")

    heading("Experience")
    body("Product Manager at Northwith Traders Inc., Jan 2020 - Mar 2022")
    body("Led roadmap for internal analytics tools.")
    body("Backend Engineer at Vantage Point Analytics, 06/2019 - 08/2021")
    body("Built ETL pipelines.")

    heading("Education")
    body("M.S. in Data Science, Innovate Institute of Technology, 2021")

    heading("Skills")
    body("Postgres, JS, K8s, React.js, TDD, Zephyr RTOS")

    c.save()


# --- 6. empty_resume.pdf / empty_resume.docx --------------------------------
def make_empty_resume_pdf() -> None:
    path = FIXTURES_DIR / "empty_resume.pdf"
    c = _pdf_canvas(path)
    c.showPage()  # a single blank page, no text/drawing operations
    c.save()


def make_empty_resume_docx() -> None:
    path = FIXTURES_DIR / "empty_resume.docx"
    d = docx.Document()
    d.save(str(path))  # a docx with no paragraphs/tables added


# --- 7. corrupt.pdf / corrupt.docx ------------------------------------------
def make_corrupt_pdf() -> None:
    # Correct PDF magic bytes so it would pass extension+magic-byte upload
    # validation, but truncated/garbage body so pdfplumber cannot parse it.
    path = FIXTURES_DIR / "corrupt.pdf"
    path.write_bytes(
        b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\nthis is not a real pdf body, just garbage bytes\n"
    )


def make_corrupt_docx() -> None:
    # Correct DOCX/zip magic bytes (PK\x03\x04) so it passes upload
    # validation, but it's missing the internal word/document.xml
    # structure python-docx requires, so opening it must fail cleanly.
    path = FIXTURES_DIR / "corrupt.docx"
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as zf:
        zf.writestr("not_a_real_docx_entry.txt", "garbage content")
    path.write_bytes(buffer.getvalue())


# --- 8. scanned_image.pdf ----------------------------------------------------
def make_scanned_image_pdf() -> None:
    path = FIXTURES_DIR / "scanned_image.pdf"
    # A plain synthetic raster image (no OCR-able text layer at all) --
    # simulates a scanned page with zero extractable text.
    img = Image.new("RGB", (400, 300), color=(200, 200, 200))
    img_path = FIXTURES_DIR / "_tmp_scan_source.jpg"
    img.save(img_path, format="JPEG")

    c = _pdf_canvas(path)
    c.drawImage(str(img_path), 72, 400, width=400, height=300)
    c.save()
    img_path.unlink()


def main() -> None:
    FIXTURES_DIR.mkdir(parents=True, exist_ok=True)
    make_normal_resume_pdf()
    make_normal_resume_docx()
    make_table_resume_docx()
    make_table_resume_pdf()
    make_two_column_resume_pdf()
    make_unusual_order_resume_pdf()
    make_headingless_resume_pdf()
    make_varied_formats_resume_pdf()
    make_empty_resume_pdf()
    make_empty_resume_docx()
    make_corrupt_pdf()
    make_corrupt_docx()
    make_scanned_image_pdf()
    print(f"Fixtures written to {FIXTURES_DIR}")


if __name__ == "__main__":
    main()
